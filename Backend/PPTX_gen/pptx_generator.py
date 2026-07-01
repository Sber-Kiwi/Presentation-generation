import json
import re

from lxml import etree
from pptx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.chart import XL_CHART_TYPE
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.text.text import _Run

EMU_PER_INCH = 914400

# Парамтеры отступов на слайде
MARGIN_X = 0.03

# Параметры заголовков (common_slide_title и title)
TITLE_Y = 0.03
TITLE_H = 0.11
TITLE_X = MARGIN_X
TITLE_W = 1.0 - (2 * MARGIN_X)
TITLE_FONT_PT = 24

# COMTIT - маленькая подпись отдела под основным заголовком
COMTIT_Y = TITLE_Y + TITLE_H
COMTIT_H = 0.04
COMTIT_X = MARGIN_X
COMTIT_W = 1.0 - (2 * MARGIN_X)
COMTIT_FONT_PT = 12

# Параметры для WaterfallData (доли от размеров слайда/блока)
WATERFALL_TOTAL_ROW_H = 0.05
WATERFALL_GAP_W = 0.015
WATERFALL_MIN_BODY_H = 0.2

# Параметры для таблицы (TABLE)
TABLE_CAPTION_H = 0.035
TABLE_CAPTION_GAP = 0.01
TABLE_NUM_COL_W = 0.03

# Именованные цвета
NAMED_COLORS = {
    "green": RGBColor(0x4C, 0xAF, 0x50),
    "red": RGBColor(0xF4, 0x43, 0x36),
    "blue": RGBColor(0x00, 0x70, 0xC0),
    "orange": RGBColor(0xFF, 0x98, 0x00),
    "gray": RGBColor(0x9E, 0x9E, 0x9E),
    "yellow": RGBColor(0xFF, 0xC1, 0x07),
    "purple": RGBColor(0x9C, 0x27, 0xB0),
    "navy": RGBColor(0x1E, 0x27, 0x61),
    "black": RGBColor(0x00, 0x00, 0x00),
    "white": RGBColor(0xFF, 0xFF, 0xFF),
}

# Цвета по умолчанию для серий (если color не распознан)
DEFAULT_SERIES_PALETTE = [
    RGBColor(0x00, 0x70, 0xC0),
    RGBColor(0xED, 0x7D, 0x31),
    RGBColor(0x4C, 0xAF, 0x50),
    RGBColor(0xFF, 0xC0, 0x00),
    RGBColor(0x9C, 0x27, 0xB0),
]

# Фиксированные цвета для Waterfall (не зависят от данных)
WATERFALL_TOTAL_COLOR = RGBColor(0x1E, 0x27, 0x61)   # START / END
WATERFALL_INC_COLOR = RGBColor(0x4C, 0xAF, 0x50)     # delta > 0
WATERFALL_DEC_COLOR = RGBColor(0xF4, 0x43, 0x36)     # delta < 0

HEADER_FILL_COLOR = RGBColor(0x00, 0x70, 0xC0)
HEADER_TEXT_COLOR = RGBColor(0xFF, 0xFF, 0xFF)


def resolve_named_color(name, fallback):
    """
    Преобразует строковое имя цвета из данных в RGBColor.
    Если имя не распознано - возвращается fallback (то что передается на вход).
    """
    if not name:
        return fallback
    return NAMED_COLORS.get(str(name).lower(), fallback)


def format_number(value):
    """123456789 -> '123 456 789' для отображения фин. значений."""
    try:
        return f"{value:,.0f}".replace(",", " ")
    except (TypeError, ValueError):
        return str(value)


class PptxBuilder:
    def __init__(self, json_path, output_path,
                 slide_width=13.333, slide_height=7.5):
        """
        Размеры слайдов указаны по умолчанию, при необходимостти
        можно указать при создании экземпляра Builder'а
        """
        self.json_path = json_path
        self.output_path = output_path
        self.prs = Presentation()
        # размер слайда по умолчанию 16:9 (13.333 : 7.5 дюйм)
        self.prs.slide_width = Inches(slide_width)
        self.prs.slide_height = Inches(slide_height)

        self.slide_w_emu = int(slide_width * EMU_PER_INCH)
        self.slide_h_emu = int(slide_height * EMU_PER_INCH)

    def _frac_rect_emu(self, x, y, width, height):
        """Конвертирует доли (0..1) слайда в абсолютные значения в EMU."""
        return (
            Emu(int(x * self.slide_w_emu)),
            Emu(int(y * self.slide_h_emu)),
            Emu(int(width * self.slide_w_emu)),
            Emu(int(height * self.slide_h_emu)),
        )

    def _get_pos_emu(self, pos):
        """Конвертирует проценты из JSON в абсолютные значения в EMU."""
        return self._frac_rect_emu(pos["x"], pos["y"], pos["width"], pos["height"])

    def _h_emu(self, frac):
        """Доля высоты слайда -> EMU (int)."""
        return int(frac * self.slide_h_emu)

    def _w_emu(self, frac):
        """Доля ширины слайда -> EMU (int)."""
        return int(frac * self.slide_w_emu)

    def build(self):
        with open(self.json_path, "r", encoding="utf-8") as f:
            data = json.load(f)

        self.global_title = data.get("meta").get("common_slide_title")

        for slide_data in data["slides"]:
            self._process_slide(slide_data)

        self.prs.save(self.output_path)

    #######################################################################
    # диспетчер объектов (создание одного слайда)
    #######################################################################

    def _process_slide(self, slide_data):
        blank_layout = self.prs.slide_layouts[6]  # пустой слайд
        slide = self.prs.slides.add_slide(blank_layout)

        if "meta" in slide_data and "title" in slide_data["meta"]:
            left, top, width, height = self._frac_rect_emu(
                TITLE_X, TITLE_Y, TITLE_W, TITLE_H
            )
            title_box = slide.shapes.add_textbox(left, top, width, height)
            tf = title_box.text_frame
            p = tf.paragraphs[0]
            p.text = slide_data["meta"]["title"]
            p.font.size = Pt(TITLE_FONT_PT)
            p.font.bold = True

        if self.global_title:
            left, top, width, height = self._frac_rect_emu(
                COMTIT_X, COMTIT_Y, COMTIT_W, COMTIT_H
            )
            comtit_box = slide.shapes.add_textbox(left, top, width, height)
            tf = comtit_box.text_frame
            p = tf.paragraphs[0]
            p.text = self.global_title
            p.font.size = Pt(COMTIT_FONT_PT)

        for obj in slide_data.get("objects", []):
            obj_type = obj["type"]
            pos = obj.get(
                "position", {"x": MARGIN_X, "y": COMTIT_Y + COMTIT_H + 0.02, "width": 1.0 - 2 * MARGIN_X, "height": 0.8}
            )
            obj_data = obj["data"]

            if obj_type == "TEXT":
                self._add_text(slide, obj_data, pos)
            elif obj_type == "TABLE":
                self._add_table(slide, obj_data, pos)
            elif obj_type == "CHART":
                self._add_chart_object(slide, obj_data, pos)

    #######################################################################
    # TEXT
    #######################################################################

    def _add_text(self, slide, data, pos):
        left, top, width, height = self._get_pos_emu(pos)

        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.word_wrap = True

        lines = data["text"].split("\n")
        
        first_para = True
        for line in lines:
            p = tf.paragraphs[0] if first_para else tf.add_paragraph()
            first_para = False
            self._apply_markdown_to_paragraph(p, line)

    def _apply_markdown_to_paragraph(self, p, line):
        """Парсит блочные элементы Markdown (заголовки, списки)."""
        default_size = Pt(14)
        
        # заголовки - # H1, ## H2, ### H3
        header_match = re.match(r'^(#{1,3})\s+(.*)', line)
        if header_match:
            level = len(header_match.group(1))
            text = header_match.group(2)
            size = Pt(28) if level == 1 else (Pt(22) if level == 2 else Pt(18))
            p.font.bold = True
            self._add_inline_runs(p, text, size)
            return

        # списки (- элемент или * элемент)
        list_match = re.match(r'^(\s*)[-*]\s+(.*)', line)
        if list_match:
            indent = len(list_match.group(1))
            text = list_match.group(2)
            level = indent // 2  # Каждые 2 пробела = новый уровень
            
            self._set_bullet(p, level)
            
            self._add_inline_runs(p, text, default_size)
            return

        # обычный текст
        self._add_inline_runs(p, line, default_size)

    def _set_bullet(self, paragraph, level=0):
        """
        Хелпер: настраивает маркер списка и уровень вложенности.
        Использует прямой доступ к XML, так как python-pptx не имеет API для списков.
        """
        pPr = paragraph._p.get_or_add_pPr()
        pPr.set('lvl', str(level))
        
        # удаление настроек маркеров
        for tag in ['a:buChar', 'a:buNone', 'a:buAutoNum']:
            for el in pPr.findall(qn(tag)):
                pPr.remove(el)
                
        # маркеры (• для первого уровня, ◦ для вложенных)
        buChar = etree.SubElement(pPr, qn('a:buChar'))
        buChar.set('char', '•' if level == 0 else '◦')

    def _add_inline_runs(self, p, text, default_size):
        """Парсит строчные элементы Markdown (**жирный**, *курсив*, `код`)."""
        # удаляем стандартный пустой run
        if len(p.runs) == 1 and p.runs[0].text == '':
            p._p.remove(p.runs[0]._r)
            
        # ***жирный курсив***, **жирный**, *курсив*, `код`
        pattern = re.compile(r'(\*\*\*(.+?)\*\*\*|\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`)')
        last_end = 0
        
        for match in pattern.finditer(text):
            # текст до фрагмента с форматированием
            if match.start() > last_end:
                run: _Run = p.add_run()
                run.text = text[last_end:match.start()]
                run.font.size = default_size
                
            run: _Run = p.add_run()
            if match.group(2):              # ***
                run.text = match.group(2)
                run.font.bold = True
                run.font.italic = True
            elif match.group(3):            # **
                run.text = match.group(3)
                run.font.bold = True
            elif match.group(4):            # *
                run.text = match.group(4)
                run.font.italic = True
            elif match.group(5):            # `
                run.text = match.group(5)
                run.font.name = 'Courier New'
                run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
                
            run.font.size = default_size
            last_end = match.end()
            
        # остаток текста
        if last_end < len(text):
            run: _Run = p.add_run()
            run.text = text[last_end:]
            run.font.size = default_size

    #######################################################################
    # TABLE
    #######################################################################

    def _add_table(self, slide, data, pos):
        headers = data["headers"]
        rows = data["rows"]
        caption = data.get("caption")


        num_rows = len(rows) + 1
        num_cols = len(headers)

        left, top, width, height = self._get_pos_emu(pos)

        # есть caption - отдаём небольшую полосу снизу
        caption_h = self._h_emu(TABLE_CAPTION_H) if caption else 0
        table_height = height - caption_h

        table_shape = slide.shapes.add_table(
            num_rows, num_cols,
            left, top, width, table_height,
        )
        table = table_shape.table

        for i, header in enumerate(headers):
            cell = table.cell(0, i)
            cell.text = str(header)
            cell.fill.solid()
            cell.fill.fore_color.rgb = HEADER_FILL_COLOR
            p = cell.text_frame.paragraphs[0]
            p.font.color.rgb = HEADER_TEXT_COLOR
            p.font.bold = True
            p.font.size = Pt(10)

        for r_idx, row in enumerate(rows):
            for c_idx, val in enumerate(row):
                cell = table.cell(r_idx + 1, c_idx)
                cell.text = str(val) if val is not None else ""
                p = cell.text_frame.paragraphs[0]
                p.font.size = Pt(10)

        if caption:
            caption_box = slide.shapes.add_textbox(
                left, top + table_height + self._h_emu(TABLE_CAPTION_GAP),
                width, caption_h,
            )
            p = caption_box.text_frame.paragraphs[0]
            p.text = caption
            p.font.size = Pt(10)
            p.font.italic = True
            p.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    #######################################################################
    # CHART (ChartData / WaterfallData)
    #######################################################################

    def _add_chart_object(self, slide, data, pos):
        kind = data["kind"]

        if kind == "ChartData":
            self._add_xy_chart(slide, data, pos)
        elif kind == "WaterfallData":
            self._add_waterfall(slide, data, pos)

    # ChartData: LINE / BAR / PIE

    def _add_xy_chart(self, slide, data, pos):
        chart_type = data.get("chart_type", "LINE")
        left, top, width, height = self._get_pos_emu(pos)

        if chart_type == "PIE":
            self._add_pie_chart(slide, data, left, top, width, height)
        elif chart_type == "BAR":
            self._add_bar_chart(slide, data, left, top, width, height)
        else:
            self._add_line_chart(slide, data, left, top, width, height)

    def _add_line_chart(self, slide, data, left, top, width, height):
        chart_data = CategoryChartData()
        categories = [point["x"] for point in data["series"][0]["data"]]
        chart_data.categories = categories

        for series in data["series"]:
            values = [point["y"] for point in series["data"]]
            chart_data.add_series(series["unit"], values)

        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.LINE_MARKERS, left, top, width, height, chart_data
        ).chart

        chart.has_legend = True
        chart.has_title = True
        chart.chart_title.text_frame.text = data.get('title', '')
        

        plot = chart.plots[0]
        plot.has_data_labels = True
        plot.data_labels.show_value = False
        plot.data_labels.font.size = Pt(9)

        for idx, series_def in enumerate(data["series"]):
            color = resolve_named_color(
                series_def.get("color"),
                DEFAULT_SERIES_PALETTE[idx % len(DEFAULT_SERIES_PALETTE)],
            )
            plot_series = plot.series[idx]
            plot_series.format.line.color.rgb = color
            plot_series.format.line.width = Pt(2.25)
            plot_series.marker.format.fill.solid()
            plot_series.marker.format.fill.fore_color.rgb = color

    def _add_bar_chart(self, slide, data, left, top, width, height):
        """
        Обычная (вертикальная) столбчатая диаграмма 
        для сравнения нескольких серий по категориям.
        """
        chart_data = CategoryChartData()
        categories = [point["x"] for point in data["series"][0]["data"]]
        chart_data.categories = categories

        for series in data["series"]:
            values = [point["y"] for point in series["data"]]
            chart_data.add_series(series["unit"], values)

        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED, left, top, width, height, chart_data
        ).chart

        chart.has_legend = True

        plot = chart.plots[0]
        plot.has_data_labels = True
        plot.data_labels.show_value = True
        plot.data_labels.font.size = Pt(9)

        for idx, series_def in enumerate(data["series"]):
            color = resolve_named_color(
                series_def.get("color"),
                DEFAULT_SERIES_PALETTE[idx % len(DEFAULT_SERIES_PALETTE)],
            )
            plot_series = plot.series[idx]
            plot_series.format.fill.solid()
            plot_series.format.fill.fore_color.rgb = color

    def _add_pie_chart(self, slide, data, left, top, width, height):
        """Круговая диаграмма строится по первой серии (категория -> доля)."""
        first_series = data["series"][0]
        chart_data = CategoryChartData()
        chart_data.categories = [point["x"] for point in first_series["data"]]
        values = [point["y"] for point in first_series["data"]]
        chart_data.add_series(first_series.get("unit", ""), values)

        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.PIE, left, top, width, height, chart_data
        ).chart

        chart.has_legend = True

        plot = chart.plots[0]
        plot.has_data_labels = True
        plot.data_labels.show_value = False
        plot.data_labels.show_percentage = True
        plot.data_labels.font.size = Pt(10)

        for i in range(len(chart_data.categories)):
            color = DEFAULT_SERIES_PALETTE[i % len(DEFAULT_SERIES_PALETTE)]
            point = plot.series[0].points[i]
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = color

    # WaterfallData: горизонтальный водопад + таблица комментариев

    def _add_waterfall(self, slide, data, pos):
        items = data["data"]

        left, top, width, height = self._get_pos_emu(pos)

        # Нижняя полоса под итоговое изменение (текст жирным)
        total_h = self._h_emu(WATERFALL_TOTAL_ROW_H)
        body_top = top
        body_height = height - total_h
        if body_height < self._h_emu(WATERFALL_MIN_BODY_H):
            # Если места мало - отдаём всю оставшуюся высоту графику/таблице.
            body_height = height
            total_h = 0

        gap = self._w_emu(WATERFALL_GAP_W)
        chart_width = int(width * 0.6)
        table_width = width - chart_width - gap

        categories, base_vals, inc_vals, dec_vals, total_vals = [], [], [], [], []
        running_total = 0
        start_delta = None
        end_delta = None

        for item in items:
            categories.append(item["name"])
            item_type = item["type"]
            plan = item["plan"]
            delta = item["delta"]

            if item_type == "START":
                base_vals.append(0)
                inc_vals.append(0)
                dec_vals.append(0)
                total_vals.append(plan)
                running_total = plan
                start_delta = delta
            elif item_type == "END":
                base_vals.append(0)
                inc_vals.append(0)
                dec_vals.append(0)
                total_vals.append(plan)
                running_total = plan
                end_delta = delta
            else:  # COMMON
                total_vals.append(0)
                if delta >= 0:
                    base_vals.append(running_total)
                    inc_vals.append(delta)
                    dec_vals.append(0)
                    running_total += delta
                else:
                    base_vals.append(running_total + delta)
                    inc_vals.append(0)
                    dec_vals.append(abs(delta))
                    running_total += delta

        chart_data = CategoryChartData()
        chart_data.categories = categories
        chart_data.add_series("Base", base_vals)
        chart_data.add_series("Increase", inc_vals)
        chart_data.add_series("Decrease", dec_vals)
        chart_data.add_series("Total", total_vals)

        graphic_frame = slide.shapes.add_chart(
            XL_CHART_TYPE.BAR_STACKED,
            left, body_top, chart_width, body_height,
            chart_data,
        )
        chart = graphic_frame.chart

        # Первый элемент данных (START) должен оказаться сверху - по
        # умолчанию горизонтальные bar-чарты рисуют первую категорию снизу.
        chart.category_axis.reverse_order = True
        chart.value_axis.has_major_gridlines = False
        chart.value_axis.tick_labels.number_format = "#,##0"
        chart.value_axis.tick_labels.number_format_is_linked = False
        chart.category_axis.tick_labels.font.size = Pt(10)
        chart.has_legend = False
        chart.has_title = False  # заголовок блока рисуется на уровне выше

        plot = chart.plots[0]
        plot.gap_width = 30

        base_series, inc_series, dec_series, total_series = plot.series

        base_series.format.fill.background()
        base_series.format.line.fill.background()

        inc_series.format.fill.solid()
        inc_series.format.fill.fore_color.rgb = WATERFALL_INC_COLOR
        inc_series.format.line.fill.background()

        dec_series.format.fill.solid()
        dec_series.format.fill.fore_color.rgb = WATERFALL_DEC_COLOR
        dec_series.format.line.fill.background()

        total_series.format.fill.solid()
        total_series.format.fill.fore_color.rgb = WATERFALL_TOTAL_COLOR
        total_series.format.line.fill.background()

        
        for s in (inc_series, dec_series, total_series):
            s.data_labels.show_value = True
            s.data_labels.font.size = Pt(10)
            s.data_labels.number_format = "#,##0"
            s.data_labels.number_format_is_linked = False

        # Жирным подписи итоговых столбцов (START / END)
        for idx, item in enumerate(items):
            if item["type"] in ("START", "END"):
                total_series.points[idx].data_label.font.bold = True

        # -- таблица комментариев справа: № + текст --
        table_left = left + chart_width + gap
        n_items = len(items)

        table_shape = slide.shapes.add_table(
            n_items, 2,
            table_left, body_top, table_width, body_height,
        )
        table = table_shape.table
        table.first_row = False
        table.first_col = False
        table.horz_banding = False
        num_col_w = self._w_emu(TABLE_NUM_COL_W)
        table.columns[0].width = Emu(num_col_w)
        table.columns[1].width = Emu(max(table_width - num_col_w, num_col_w))

        row_height = Emu(int(body_height / n_items))
        for row in table.rows:
            row.height = row_height

        for i, item in enumerate(items):
            is_total_row = item["type"] in ("START", "END")
            fill_color = RGBColor(0xE7, 0xEC, 0xF5) if is_total_row else RGBColor(0xFF, 0xFF, 0xFF)

            num_cell = table.cell(i, 0)
            num_cell.text = str(i + 1)
            num_cell.fill.solid()
            num_cell.fill.fore_color.rgb = fill_color
            num_p = num_cell.text_frame.paragraphs[0]
            num_p.font.size = Pt(10)
            num_p.font.bold = is_total_row
            num_p.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

            text_cell = table.cell(i, 1)
            text_cell.text = (item.get("text") or "").strip()
            text_cell.fill.solid()
            text_cell.fill.fore_color.rgb = fill_color
            text_p = text_cell.text_frame.paragraphs[0]
            text_p.font.size = Pt(10)
            text_p.font.bold = is_total_row
            text_p.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

        # -- итоговая строка с общим изменением (delta START == delta END) --
        if total_h > 0:
            total_delta = start_delta if start_delta is not None else end_delta
            if total_delta is not None:
                sign = "+" if total_delta >= 0 else ""
                total_text = f"Итоговое изменение: {sign}{format_number(total_delta)}"

                total_box = slide.shapes.add_textbox(
                    left, body_top + body_height + self._h_emu(TABLE_CAPTION_GAP),
                    width, total_h,
                )
                p = total_box.text_frame.paragraphs[0]
                p.text = total_text
                p.font.size = Pt(14)
                p.font.bold = True


if __name__ == "__main__":
    builder = PptxBuilder("test_data.json", "output.pptx")
    builder.build()
